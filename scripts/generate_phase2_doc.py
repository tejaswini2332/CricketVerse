import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

NAVY = RGBColor(15, 23, 42)
GREEN = RGBColor(21, 128, 61)
MUTED = RGBColor(100, 116, 139)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# Header
title = doc.add_paragraph()
title_run = title.add_run("CRICKETVERSE: A Decentralized Web3 Cricket Ecosystem")
title_run.font.name = "Arial"
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.color.rgb = NAVY
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph()
sub_run = sub.add_run("Phase 2 Progress Report: Completed Days 11–20 & Final Sprint (Days 21–30) Workplan")
sub_run.font.name = "Arial"
sub_run.font.size = Pt(13)
sub_run.font.bold = True
sub_run.font.color.rgb = GREEN
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(12)
m_run = meta.add_run(
    "Author: Tejaswini Kunkulagunta | Roll No: CS23B032 | Dept: CSE | Academic Year: 2026\n"
    "Phase 2 Stack: Chainlink Functions, OpenZeppelin Escrow, On-Chain SVG (EVM Cancun), Dockerized PostgreSQL & Redis"
)
m_run.font.name = "Arial"
m_run.font.size = Pt(9.5)
m_run.font.italic = True
m_run.font.color.rgb = MUTED
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

def add_heading(text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = NAVY
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = GREEN
    return h

# 1. Executive Summary
add_heading("1. Executive Summary (Sprint 2 / Days 11–20)", 1)
doc.add_paragraph(
    "During Phase 2 (Days 11–20), the CricketVerse ecosystem transitioned from initial smart contract prototypes "
    "into an automated, high-throughput decentralized platform. Engineering objectives focused on: (1) trustless sports "
    "oracle consensus via Chainlink Functions architecture, (2) non-custodial tournament prize pool escrow smart contracts "
    "with automated winner disbursements, (3) pure on-chain dynamic SVG vector graphics rendering athletic milestone badges, "
    "and (4) a Dockerized hybrid database tier utilizing PostgreSQL relational tables and Redis sub-second caching.\n\n"
    "The test suite was expanded from 13 to 27 unit tests, all executing in 860ms with a 100% pass rate under Hardhat's "
    "Cancun EVM target. All contract deployments and interactive Web3 DApp views have been fully validated."
)

# 2. Phase 2 Completed Engineering Deliverables
add_heading("2. Phase 2 Completed Engineering Deliverables", 1)

deliverables_data = [
    ("Days 11–12", "Chainlink Functions Oracle Integration", 
     "• CricketOracleConsumer.sol: Authored Chainlink Functions consumer contract.\n"
     "• Bridges real-world sports APIs (CricAPI/Sportradar) into the EVM.\n"
     "• Emits StatsRequested and securely verifies relayer signatures before mutating PlayerCardNFT.\n"
     "• Prevents duplicate request fulfillment and supports batch post-match settlement."),
    
    ("Days 13–14", "Dockerized Hybrid Off-Chain Infrastructure", 
     "• docker-compose.yml: Orchestrates PostgreSQL 16 Alpine, Redis 7 Alpine, and Node.js REST API.\n"
     "• server/schema.sql: 6 relational tables (match_fixtures, player_telemetry, ball_telemetry, fantasy_squads, tournament_contests, oracle_audit_logs).\n"
     "• server/db.js: High-performance connector providing sub-second Redis leaderboard caching with in-memory test fallbacks."),
     
    ("Days 15–16", "On-Chain Dynamic SVG Card Generator", 
     "• CardSVGGenerator.sol: Computes responsive vector graphics on-chain encoded in RFC-2397 base64 Data URIs.\n"
     "• Dynamic Milestones: Evaluates athletic criteria to unlock visual badges:\n"
     "  - CENTURION 100+ (Runs >= 100)\n"
     "  - HALF-TON 50+ (Runs >= 50)\n"
     "  - 5-WKT HAUL (Wickets >= 5)\n"
     "  - STRIKE MASTER (Strike Rate >= 140.00)\n"
     "  - ECONOMY KING (Economy <= 6.50)\n"
     "• Rarity Color Coding: Gold (#F59E0B) for Legendary down to Slate (#64748B) for Common."),
     
    ("Days 17–18", "Tournament Escrow Prize Pool Smart Contract", 
     "• TournamentPool.sol: Non-custodial fantasy tournament manager locking CRIC entry fees in escrow.\n"
     "• 5% Platform Treasury Fee: Automatically routes 500 bps to organizer/liquidity treasury.\n"
     "• 95% Net Winner Distribution: Trustless pro-rata settlement to verified leaderboard winners.\n"
     "• Rain / Abandonment Protection: Full 100% refund claimable if match is washed out.\n"
     "• Security: OpenZeppelin ReentrancyGuard and Checks-Effects-Interactions pattern."),
     
    ("Days 19–20", "Automated Testing, Deployment Automation & DApp Integration", 
     "• test/Phase2_Ecosystem.test.js: 14 new comprehensive unit tests added (total 27 passing tests).\n"
     "• scripts/deploy_phase2.js: Automated multi-contract deployment script for Hardhat & Polygon Amoy.\n"
     "• Web3 Frontend: Added Tournament Arena and Chainlink & Dynamic SVG Live Arena tabs to public DApp.")
]

table_deliv = doc.add_table(rows=1, cols=3)
table_deliv.alignment = WD_TABLE_ALIGNMENT.CENTER
t_hdr = table_deliv.rows[0].cells
t_hdr[0].text = "Sprint Window"
t_hdr[1].text = "Engineering Milestone"
t_hdr[2].text = "Technical Deliverables"
for c in t_hdr:
    set_cell_background(c, "15803D")
    for r in c.paragraphs[0].runs:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

for dw, em, td in deliverables_data:
    row = table_deliv.add_row().cells
    row[0].text = dw
    row[1].text = em
    row[2].text = td

# 3. Automated Test Suite Metrics
add_heading("3. Automated Test Suite Execution Metrics", 1)
doc.add_paragraph(
    "Both Phase 1 and Phase 2 test suites were executed sequentially on the Cancun EVM target with optimizer enabled (200 runs, viaIR: true):\n"
    "• Total Tests: 27 passing (0 failed, 0 skipped)\n"
    "• Total Execution Duration: 860 ms\n"
    "• Test Coverage Scope: ERC-20 faucet & supply bounds, ERC-1155 rarity caps & stats mutation, "
    "ERC-721 15% resale markup ceiling, turnstile single-entry anti-scalping, non-custodial marketplace swap, "
    "Chainlink oracle request/fulfill lifecycle, TournamentPool CRIC escrow locking, 5% treasury routing, "
    "95% winner distribution, 100% rain refund, and on-chain SVG badge triggers."
)

# 4. Visual Verification & Frontend Showcase
add_heading("4. Visual Verification: Phase 2 User Interface Showcase", 1)
doc.add_paragraph("The updated DApp running on http://localhost:3000 highlights the new Phase 2 Web3 interfaces:")

screenshots = [
    ("4.1 Tournament Arena: Escrow Staking & Sub-Second Leaderboard",
     "/Users/tejaswinikunkulagunta/.gemini/antigravity-ide/brain/f60a7274-063d-4d77-9664-ba5d621f8799/tournament_arena_1792040191861.png",
     "Figure 1: Tournament Arena displaying active CRIC prize pools (100,000 CRIC Grand League), live entry staking (50 CRIC), and sub-second Redis leaderboard cache."),
     
    ("4.2 Non-Custodial Contest Participation & Purse Deduction",
     "/Users/tejaswinikunkulagunta/.gemini/antigravity-ide/brain/f60a7274-063d-4d77-9664-ba5d621f8799/joined_contest_1792040214357.png",
     "Figure 2: Real-time confirmation of tournament entry fee transfer into the smart contract escrow, updating participant capacity and purse balances."),
     
    ("4.3 Chainlink Functions Oracle Simulator & Dynamic SVG Studio",
     "/Users/tejaswinikunkulagunta/.gemini/antigravity-ide/brain/f60a7274-063d-4d77-9664-ba5d621f8799/chainlink_tab_1792040236833.png",
     "Figure 3: Chainlink Functions oracle monitoring portal connected to decentralized relayer, featuring pure on-chain dynamic SVG vector card rendering."),
     
    ("4.4 Real-Time Oracle Telemetry Finalization & On-Chain Badge Mutation",
     "/Users/tejaswinikunkulagunta/.gemini/antigravity-ide/brain/f60a7274-063d-4d77-9664-ba5d621f8799/oracle_synced_1792040265660.png",
     "Figure 4: Automated match finalization showing cryptographic audit logs, updated career runs/wickets, and dynamically earned milestone achievement badges (Centurion 100+, Strike Master).")
]

for title_text, img_path, caption in screenshots:
    add_heading(title_text, 2)
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6.4))
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(8)
        c_run = cap.add_run(caption)
        c_run.font.name = "Arial"
        c_run.font.size = Pt(8.5)
        c_run.font.italic = True
        c_run.font.color.rgb = MUTED

# 5. Workplan for Phase 3: Final Sprint (Days 21–30)
add_heading("5. Final Sprint Workplan (Days 21–30: Testnet & Capstone Defense)", 1)
doc.add_paragraph(
    "With Phase 1 and Phase 2 successfully completed, the project enters its final 10-day sprint towards public testnet "
    "deployment, rigorous static analysis security audits, mobile gatekeeper client integration, and capstone thesis defense:"
)

table_p3 = doc.add_table(rows=1, cols=3)
table_p3.alignment = WD_TABLE_ALIGNMENT.CENTER
t_p3_hdr = table_p3.rows[0].cells
t_p3_hdr[0].text = "Timeline"
t_p3_hdr[1].text = "Milestone Objectives"
t_p3_hdr[2].text = "Target Deliverables"
for c in t_p3_hdr:
    set_cell_background(c, "15803D")
    for r in c.paragraphs[0].runs:
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

p3_rows = [
    ("Days 21–23", "Polygon Amoy Testnet Deployment & Etherscan Verification", 
     "Acquire Amoy MATIC from official faucet; deploy CRIC, PlayerCardNFT, MatchTicketNFT, Marketplace, CardSVGGenerator, and TournamentPool to Polygon Amoy; verify contracts on Polygonscan."),
    ("Days 24–25", "Static Analysis (Slither) & Gas Optimization", 
     "Execute Slither security audit; profile gas consumption per transaction (listing, buying, turnstile check-in) using Hardhat Gas Reporter; optimize storage packing."),
    ("Days 26–27", "Mobile Turnstile Gatekeeper PWA", 
     "Build lightweight Progressive Web App (PWA) with HTML5 Camera QR scanner for offline-tolerant gatekeeper admission control at stadium turnstiles."),
    ("Days 28–30", "Academic Project Thesis, Defense Deck & Final Submission", 
     "Compile final B.Tech capstone thesis document; prepare comparative Web2 vs Web3 performance benchmarks; create video walkthrough and presentation slides.")
]

for nr in p3_rows:
    row = table_p3.add_row().cells
    row[0].text = nr[0]
    row[1].text = nr[1]
    row[2].text = nr[2]

output_filename = "CricketVerse_Phase2_Progress_Report.docx"
doc.save(output_filename)
print(f"Generated: {output_filename} ({os.path.getsize(output_filename)} bytes)")
