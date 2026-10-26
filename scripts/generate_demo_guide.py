from pathlib import Path
from PIL import Image, ImageFilter, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'demo_assets'
OUT = ROOT / 'CricketVerse_Demo_Guide.pptx'
W, H = 13.333, 7.5
BG = '071A2D'; PANEL = '102A42'; PANEL2 = '153650'; WHITE = 'F4F8FC'
MUTED = 'A9BDCE'; TEAL = '33D6B0'; GOLD = 'F7C55A'; RED = 'FF7B78'; BLUE = '67B7FF'

prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

def rgb(h): return RGBColor.from_string(h)
def shape(slide, x, y, w, h, color=PANEL, radius=True, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color); sh.line.fill.background() if not line else None
    if line: sh.line.color.rgb = rgb(line)
    return sh
def text(slide, x, y, w, h, value, size=16, color=WHITE, bold=False, font='Aptos', align=None, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(.04); tf.margin_right = Inches(.04); tf.margin_top = Inches(.02); tf.margin_bottom = Inches(.02); tf.vertical_anchor = valign
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = line; p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color)
        if align is not None: p.alignment = align
    return box
def base(title, kicker='CRICKETVERSE  /  DEMO FIELD GUIDE', num=None):
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    text(s,.55,.28,10,.28,kicker,10,TEAL,True)
    text(s,.55,.72,12,.6,title,28,WHITE,True)
    shape(s,.55,1.43,12.2,.018,PANEL2,False)
    if num: text(s,12.1,7.08,.65,.2,f'{num:02}',10,MUTED,True,align=PP_ALIGN.RIGHT)
    return s
def pill(s,x,y,label,color=TEAL,w=1.35):
    shape(s,x,y,w,.32,color); text(s,x+.06,y+.045,w-.12,.21,label,9,BG,True,align=PP_ALIGN.CENTER)
def image_fit(s,path,x,y,w,h):
    im=Image.open(path).convert('RGB'); target=(1600,1000); ratio=w/h; ir=im.width/im.height
    if ir > ratio:
        nw=int(im.height*ratio); left=(im.width-nw)//2; im=im.crop((left,0,left+nw,im.height))
    else:
        nh=int(im.width/ratio); top=(im.height-nh)//2; im=im.crop((0,top,im.width,top+nh))
    temp=ASSETS/'deck-crop.jpg'; im.save(temp,quality=88)
    s.shapes.add_picture(str(temp), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
def label(s,x,y,w,k,v,accent=TEAL):
    shape(s,x,y,w,1.1,PANEL)
    text(s,x+.16,y+.12,w-.3,.24,k.upper(),9,accent,True)
    text(s,x+.16,y+.43,w-.3,.52,v,15,WHITE,True)
def flow(s,items,y=2.15):
    n=len(items); gap=.34; w=(12.1-(n-1)*gap)/n; x=.6
    for i,(head,sub) in enumerate(items):
        shape(s,x,y,w,1.3,PANEL)
        text(s,x+.16,y+.15,w-.32,.28,head,14,TEAL,True)
        text(s,x+.16,y+.54,w-.32,.6,sub,11,MUTED)
        if i<n-1: text(s,x+w+.04,y+.48,gap-.08,.3,'→',20,GOLD,True,align=PP_ALIGN.CENTER)
        x += w+gap

# cover
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
shape(s,8.3,0,5.03,7.5,PANEL,False)
text(s,.7,.55,5,.3,'CRICKETVERSE  /  BLOCKCHAIN DEMO',11,TEAL,True)
text(s,.7,1.35,7.3,1.5,'Own the moment.\nProve the play.',38,WHITE,True)
text(s,.75,3.25,6.6,.85,'A hands-on guide to wallets, collectibles, trading, tournaments, fan polls and ticket entry.',18,MUTED)
pill(s,.75,4.55,'LOCAL DEMO',TEAL,1.35); pill(s,2.25,4.55,'EVM · HARDHAT',GOLD,1.7)
image_fit(s,ASSETS/'ui-auction.png',8.65,.55,4.25,4.0)
text(s,8.75,4.85,3.9,.45,'Real CricketVerse UI · Auction Room',12,WHITE,True)
text(s,.75,6.8,7,.3,'Classroom prototype  •  chain ID 31337  •  seeded demo state',11,MUTED)

# 2 quickstart
s=base('Get ready in three steps',num=2)
flow(s,[('01  START CHAIN','Terminal 1\nnpm run chain:node'),('02  START APP','Terminal 2\nnpm start'),('03  CONNECT','Chrome + MetaMask\nhttp://localhost:3000')])
shape(s,.65,4.0,5.85,2.2,PANEL)
text(s,.9,4.25,5.35,.3,'WALLET SETUP',12,GOLD,True)
text(s,.9,4.72,5.2,1.1,'Open in Chrome with MetaMask. Select CricketVerse Local Hardhat, connect, then click Get 20 demo ETH for gas.',16,WHITE)
shape(s,6.8,4.0,5.85,2.2,PANEL)
text(s,7.05,4.25,5.3,.3,'IMPORTANT',12,RED,True)
text(s,7.05,4.72,5.2,1.1,'Seeded player wallets start with 10,000 CRIC. Fresh wallets can claim 1,000 CRIC; native ETH pays gas.',16,WHITE)

# 3 architecture
s=base('One wallet. Contracts do the settling.',num=3)
flow(s,[('PLAYER','MetaMask\nEIP-1193 signer'),('WEB APP','ethers.js\nreads + transactions'),('LOCAL EVM','Hardhat RPC\nchain 31337'),('CONTRACTS','CRIC · NFTs ·\nmarket · polls')],2.0)
shape(s,.75,4.15,11.8,1.65,PANEL)
text(s,1.0,4.4,11.2,.3,'WHAT “ON-CHAIN” MEANS HERE',12,TEAL,True)
text(s,1.0,4.85,11.15,.7,'Wallet signs a transaction → contract checks rules → EVM updates state → receipt and events can be inspected. The browser is a client; ownership and settlement live in deployed contracts.',16,WHITE)
text(s,.75,6.25,11,.3,'Local Hardhat is a classroom chain on this computer. Shared access requires a public testnet deployment and hosted site.',11,MUTED)

# 4 wallet and faucet
s=base('Wallet + 1,000 CRIC faucet',num=4)
image_fit(s,ASSETS/'ui-collectibles.png',.65,1.8,6.0,4.85)
shape(s,6.95,1.85,5.6,1.25,PANEL); text(s,7.2,2.05,5.1,.25,'CLAIM FLOW',11,TEAL,True); text(s,7.2,2.42,5.0,.48,'Connect → check gas → claimFaucet() → 1,000 CRIC minted',15,WHITE,True)
shape(s,6.95,3.35,5.6,1.25,PANEL); text(s,7.2,3.55,5.1,.25,'WHY YOUR CLAIM FAILED',11,RED,True); text(s,7.2,3.92,5.0,.5,'Connected wallet had 0 Hardhat ETH. The contract cannot pay gas on your behalf.',14,WHITE)
shape(s,6.95,4.85,5.6,1.25,PANEL); text(s,7.2,5.05,5.1,.25,'LOCAL DEMO FIX',11,GOLD,True); text(s,7.2,5.42,5.0,.5,'Click Get 20 demo ETH, then claim 1,000 CRIC in this wallet.',14,WHITE)
text(s,.75,6.83,11,.25,'Screenshot: collectibles UI. The wallet address is intentionally cropped out in this deck.',10,MUTED)

# 5 cards
s=base('Collectible cards: mint, own, transfer',num=5)
image_fit(s,ASSETS/'ui-collectibles.png',.65,1.85,5.25,4.65)
flow(s,[('REGISTERED','Catalog\nentry'),('MINT','PlayerCardNFT\nERC-1155'),('OWN','Wallet balance\non-chain')],2.1)
shape(s,6.2,4.1,6.3,1.8,PANEL)
text(s,6.45,4.35,5.8,.27,'DEMO SCRIPT',11,GOLD,True)
text(s,6.45,4.78,5.75,.9,'Claim CRIC → collect a card → confirm MetaMask → inspect the card balance in My Collection.',15,WHITE)

# 6 marketplace
s=base('Auction Room: fixed-price card + ticket trades',num=6)
image_fit(s,ASSETS/'ui-auction.png',.65,1.85,6.0,4.8)
shape(s,6.95,1.9,5.55,1.1,PANEL); text(s,7.2,2.08,5.1,.24,'SEEDED LIVE LISTINGS',11,TEAL,True); text(s,7.2,2.43,5.0,.36,'Card · 125 CRIC     Ticket · 90 CRIC',15,WHITE,True)
flow(s,[('APPROVE','CRIC allowance'),('BUY','Marketplace\nchecks + transfers'),('RECEIPT','Ownership +\nsettlement events')],3.45)
text(s,7.0,5.65,5.3,.6,'Buy from a wallet other than the seller. These are fixed-price listings, not timed bidding auctions.',12,MUTED)

# 7 tickets
s=base('Ticketing: collectible proof, controlled entry',num=7)
flow(s,[('HOLD','MatchTicketNFT\nERC-721'),('SIGN','Holder signs\nentry pass'),('SCAN','Gatekeeper\nverifies'),('REDEEM','Contract marks\nused once')],2.0)
shape(s,.75,4.2,7.0,1.6,PANEL); text(s,1.0,4.45,6.5,.25,'CONTRACT CHECKS',11,TEAL,True); text(s,1.0,4.88,6.4,.65,'Correct ticket owner · authorized gatekeeper · fixture time window · not previously redeemed',16,WHITE)
shape(s,8.05,4.2,4.45,1.6,PANEL); text(s,8.3,4.45,4,.25,'DEMO NOTE',11,GOLD,True); text(s,8.3,4.88,4,.65,'Use a wallet with GATEKEEPER_ROLE for scan verification.',14,WHITE)

# 8 tournaments
s=base('Tournament: entry escrow → organizer settlement',num=8)
flow(s,[('BUILD XI','Squad committed\nas hash'),('JOIN','50 CRIC entry\ninto escrow'),('LOCK','Organizer closes\nentry'),('SETTLE','PrizePaid events\nto winners')],2.0)
shape(s,.75,4.2,7.2,1.6,PANEL); text(s,1.0,4.45,6.7,.25,'SEEDED PAST SETTLEMENT',11,TEAL,True); text(s,1.0,4.88,6.5,.65,'100 CRIC pooled → 5% fee → 66.5 + 28.5 CRIC prizes (two PrizePaid events)',15,WHITE)
shape(s,8.25,4.2,4.25,1.6,PANEL); text(s,8.5,4.45,3.8,.25,'TRUTH IN DEMO',11,RED,True); text(s,8.5,4.88,3.8,.65,'Organizer inputs winners; no live score oracle calculates results.',13,WHITE)

# 9 polls
s=base('Fan polls: one wallet, one recorded vote',num=9)
shape(s,.75,1.95,5.5,3.8,PANEL); text(s,1.0,2.2,5,.3,'ON-CHAIN VOTE PATH',12,TEAL,True)
text(s,1.0,2.8,4.9,2.4,'Choose an option\n↓\nSign vote transaction\n↓\nFanPolls records voter + option\n↓\nVote count increments',18,WHITE,True)
shape(s,6.7,1.95,5.8,1.55,PANEL); text(s,6.95,2.18,5.25,.25,'SEEDED POLLS',11,GOLD,True); text(s,6.95,2.58,5.2,.62,'Match winner · wicket range\nOpen; vote counts start at zero',15,WHITE)
shape(s,6.7,3.85,5.8,1.9,PANEL); text(s,6.95,4.08,5.25,.25,'RULE',11,TEAL,True); text(s,6.95,4.48,5.2,.95,'One address can vote once per poll. A vote is not a CRIC reward, and seeded poll questions are classroom data.',15,WHITE)

# 10 UI screenshots
s=base('The experience: collectible shelf + trading floor',num=10)
image_fit(s,ASSETS/'ui-collection.png',.65,1.85,5.75,4.75)
image_fit(s,ASSETS/'ui-auction.png',6.8,1.85,5.75,4.75)
text(s,.8,6.72,5.4,.25,'MY COLLECTION · screenshot captured with an empty personal wallet',10,MUTED)
text(s,6.95,6.72,5.4,.25,'AUCTION ROOM · seeded listings visible',10,MUTED)

# 11 proof
s=base('How to prove a feature really used the chain',num=11)
flow(s,[('1 · WALLET','Confirm account +\nchain ID'),('2 · TX HASH','Open transaction\nin local explorer / logs'),('3 · EVENT','Find event emitted\nby expected contract'),('4 · STATE','Read balance,\nowner or status')],2.0)
shape(s,.75,4.2,11.8,1.55,PANEL); text(s,1.0,4.44,11.2,.26,'GOOD DEMO EVIDENCE',11,TEAL,True); text(s,1.0,4.85,11.1,.55,'Show MetaMask confirmation, the mined receipt, emitted event, and resulting state. For ticket scan and polls, use a fresh test wallet so the one-time rule is visible.',15,WHITE)

# 12 limitations
s=base('What is live, seeded, or still a prototype',num=12)
label(s,.7,1.85,3.85,'LIVE CONTRACT LOGIC','Token · cards · tickets\nmarket · pools · polls',TEAL)
label(s,4.75,1.85,3.85,'SEEDED CHAIN DATA','Listings · fixtures · polls\nthree contest states',GOLD)
label(s,8.8,1.85,3.85,'NOT LIVE DATA','Player catalog · fantasy\npoints · cricket oracle',RED)
shape(s,.75,3.35,11.8,2.0,PANEL); text(s,1.0,3.62,11.2,.3,'CLASSROOM LIMITS TO SAY OUT LOUD',12,GOLD,True); text(s,1.0,4.08,11.0,1.0,'Local chain only · no public deployment · no live scores/oracle · organizer chooses winners · auction is fixed price · UI screenshots show seeded interface, not proof that every user action was executed during capture.',16,WHITE)
text(s,.8,6.2,11,.35,'The code and tests verify contract behavior; the presentation should distinguish reproducible actions from seeded example state.',11,MUTED)

# 13 walkthrough
s=base('Suggested 6-minute live walkthrough',num=13)
flow(s,[('0:00','Connect MetaMask\nwallet'),('0:30','Get local ETH\nfor gas'),('1:00','Claim CRIC +\ncollect a card'),('2:00','Buy seeded card\nfrom another wallet'),('3:30','Join contest +\nshow settlement'),('5:00','Vote + scan\nwith gatekeeper')],2.0)
shape(s,.75,4.25,11.8,1.35,PANEL); text(s,1.0,4.5,11.2,.25,'BEFORE YOU PRESENT',11,TEAL,True); text(s,1.0,4.92,11,.45,'For each wallet: connect → Get 20 demo ETH → claim 1,000 CRIC. Seeded player wallets start at 10,000 CRIC.',15,WHITE)
text(s,.8,6.15,11.5,.4,'Polls record votes but pay no CRIC. Card/ticket buys spend CRIC; tournament rewards require organizer settlement.',11,MUTED)

prs.save(OUT)
print(f'Wrote {OUT} ({len(prs.slides)} slides)')
